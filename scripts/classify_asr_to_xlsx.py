#!/usr/bin/env python3
"""工序一后置统一分类分拣入口：六路 ASR 已齐 → 五类分拣 → 完整 XLSX。

默认（035 / 034）：五类 v2.3 最小声学证据
  audio_energy → silero_vad → efficientat_candidates → efficientat → classify
  产物 stem：*_five_class_v2_3_min_acoustic_*
  与 pipelines/classify_dataset_five_class_v2_3_min_acoustic.yaml 语义等价。

兼容旧批：--rule-version v2_2_auto_noise
  audio_energy → dnsmos_v2_2_candidates → classify(v2.2)
  产物 stem：*_five_class_v2_2_auto_noise_*（隔离，不覆盖 v2.3 默认路径）

不启动 ASR、不伪造 register/reservation/release。

别名解析顺序：
  1. 显式 --family qwen=run1,run2（可重复三次）
  2. 批次 dataset YAML 的 model_families
  3. stage1 固定别名 qwen_1/2、glm_1/2、sensevoice_1/2

示例（默认 v2.3）::

  python scripts/classify_asr_to_xlsx.py --batch \"$BATCH\" --energy-workers 32

显式 legacy v2.2::

  python scripts/classify_asr_to_xlsx.py --batch \"$BATCH\" \\
    --rule-version v2_2_auto_noise --energy-workers 32
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import yaml

from audio_engine.core.manifest import Manifest
from audio_engine.core.operator import OperatorConfig
from audio_engine.core.selection_v3.config import SelectionV3Config
from audio_engine.core.selection_v3.input_contract import (
    original_audio_sha256,
    validate_base_snapshot,
)
from audio_engine.core.selection_v3.types import (
    RULE_VERSION_FIVE_CLASS_V2_2,
    RULE_VERSION_FIVE_CLASS_V2_3,
)
from audio_engine.core.source_naming import validate_asr_run, validate_source_name
from audio_engine.core.stage1.cache_policy import FAMILY_RUN_ALIASES, REQUIRED_FAMILIES
from audio_engine.operators.quality.aggregate_manifests import AggregateManifestsOperator
from audio_engine.operators.quality.audio_energy import AudioEnergyOperator
from audio_engine.operators.quality.classify import ClassifyOperator
from audio_engine.operators.quality.dnsmos_v2_2_candidates import DnsmosV22CandidatesOperator
from audio_engine.operators.quality.efficientat import EfficientAtOperator
from audio_engine.operators.quality.efficientat_candidates import EfficientAtCandidatesOperator
from audio_engine.operators.quality.silero_vad import SileroVadOperator

CATEGORIES = {"voicemail", "semantic_risk", "environment_noise", "gold_candidate", "hardcase"}

# CLI short names (035); map to selection rule_version / stems / configs.
RULE_CLI_V23 = "v2_3_min_acoustic"
RULE_CLI_V22 = "v2_2_auto_noise"
DEFAULT_RULE_VERSION = RULE_CLI_V23

RULE_PROFILES = {
    RULE_CLI_V23: {
        "selection_rule": RULE_VERSION_FIVE_CLASS_V2_3,
        "selection_config": "configs/selection/zh_asr_five_class_v2_3_min_acoustic.yaml",
        "xlsx_stem": "summary_five_class_v2_3_min_acoustic",
        "classified_stem": "classified_five_class_v2_3_min_acoustic",
        "pipeline_yaml": "pipelines/classify_dataset_five_class_v2_3_min_acoustic.yaml",
        "engine": "v2_3",
    },
    RULE_CLI_V22: {
        "selection_rule": RULE_VERSION_FIVE_CLASS_V2_2,
        "selection_config": "configs/selection/zh_asr_five_class_v2_2_auto_noise.yaml",
        "xlsx_stem": "summary_five_class_v2_2_auto_noise",
        "classified_stem": "classified_five_class_v2_2_auto_noise",
        "pipeline_yaml": "pipelines/classify_dataset_five_class_v2_2_auto_noise.yaml",
        "engine": "v2_2",
    },
}

V23_QUALITY_CONFIGS = (
    "configs/quality/silero_vad_v1.yaml",
    "configs/quality/efficientat_v1.yaml",
    "configs/quality/acoustic_evidence_v1.yaml",
)


@dataclass(frozen=True)
class RouteCheck:
    family: str
    alias: str
    path: Path
    exists: bool


class PreflightError(ValueError):
    """Structured preflight failure with per-route inventory."""

    def __init__(self, batch: str, message: str, routes: list[RouteCheck] | None = None):
        super().__init__(message)
        self.batch = batch
        self.routes = list(routes or [])


def load_yaml(path):
    value = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise TypeError(f"配置必须是 YAML 映射: {path}")
    return value


def resolve_rule_profile(rule_version: str) -> dict:
    key = str(rule_version or DEFAULT_RULE_VERSION).strip()
    if key not in RULE_PROFILES:
        raise ValueError(
            f"不支持的 --rule-version={rule_version!r}；"
            f"可选: {RULE_CLI_V23}（默认）, {RULE_CLI_V22}（legacy）"
        )
    return dict(RULE_PROFILES[key])


def assert_v23_dependencies_ready() -> None:
    """Fail loudly if 034 artifacts are missing; never silently fall back to DNSMOS."""
    missing: list[str] = []
    profile = RULE_PROFILES[RULE_CLI_V23]
    for rel in (
        profile["pipeline_yaml"],
        profile["selection_config"],
        *V23_QUALITY_CONFIGS,
    ):
        path = ROOT / rel
        if not path.is_file():
            missing.append(rel)
    try:
        from audio_engine.operators.quality import (  # noqa: F401
            efficientat,
            efficientat_candidates,
            silero_vad,
        )
        from audio_engine.core.selection_v3.classifier import classify_sample  # noqa: F401
        from audio_engine.core.selection_v3.types import is_five_class_v2_3_rule

        if not is_five_class_v2_3_rule(RULE_VERSION_FIVE_CLASS_V2_3):
            missing.append("selection_five_class_v2_3_min_acoustic classifier hook")
    except ImportError as exc:
        missing.append(f"import: {exc}")
    if missing:
        raise FileNotFoundError(
            "v2.3 入口依赖未就绪（须先落地 034，或同批交付引擎）。缺失：\n  - "
            + "\n  - ".join(missing)
            + f"\n可改用 --rule-version {RULE_CLI_V22} 复现旧口径；禁止用 DNSMOS 冒充 v2.3。"
        )


def default_dataset_config_path(batch: str) -> Path:
    return ROOT / "configs/datasets" / f"zh_asr_v3_{batch.replace('-', '_')}.yaml"


def stage1_default_families() -> dict[str, list[str]]:
    return {family: list(FAMILY_RUN_ALIASES[family]) for family in REQUIRED_FAMILIES}


def _normalize_families(raw: dict) -> dict[str, list[str]]:
    if not isinstance(raw, dict):
        raise ValueError("model_families 必须是映射")
    families: dict[str, list[str]] = {}
    for family, aliases in raw.items():
        if not isinstance(aliases, (list, tuple)) or len(aliases) != 2:
            raise ValueError(f"家族 {family} 须恰好两个 ASR 别名")
        families[str(family)] = [validate_asr_run(x) for x in aliases]
    if len(families) != 3:
        raise ValueError("需要三个模型家族，每族恰好两个 ASR 别名")
    if "qwen" not in families:
        raise ValueError("当前规则模板以 qwen 为 target_family，三族中须包含 qwen")
    flat = [alias for aliases in families.values() for alias in aliases]
    if len(set(flat)) != 6:
        raise ValueError("六路 ASR 别名必须互异")
    return families


def resolve_families(args) -> tuple[dict[str, list[str]], str]:
    """Return (families, resolution_source). Priority: CLI > dataset YAML > stage1 defaults."""
    if args.family:
        families: dict[str, list[str]] = {}
        for item in args.family:
            family, sep, aliases = item.partition("=")
            if not sep or family in families:
                raise ValueError("--family 格式为 qwen=run1,run2，家族不可重复")
            families[family] = [validate_asr_run(x) for x in aliases.split(",") if x.strip()]
        return _normalize_families(families), "cli"

    if args.dataset_config is not None:
        path = Path(args.dataset_config)
        if not path.is_file():
            raise ValueError(f"批次 dataset 配置不存在: {path}")
        families = _normalize_families(load_yaml(path).get("model_families", {}))
        return families, f"dataset:{path}"

    auto = default_dataset_config_path(args.batch)
    if auto.is_file():
        families = _normalize_families(load_yaml(auto).get("model_families", {}))
        return families, f"dataset:{auto}"

    return stage1_default_families(), "stage1_default"


def format_route_inventory(routes: list[RouteCheck]) -> list[str]:
    width = max((len(r.alias) for r in routes), default=8)
    lines = []
    for route in routes:
        mark = "OK" if route.exists else "缺失"
        lines.append(
            f"  [{route.family}] {route.alias:<{width}} {mark:<4} {route.path.as_posix()}"
        )
    return lines


def summarize_missing(routes: list[RouteCheck]) -> str:
    by_family: dict[str, list[RouteCheck]] = {}
    for route in routes:
        by_family.setdefault(route.family, []).append(route)
    parts: list[str] = []
    for family, items in by_family.items():
        missing = [r.alias for r in items if not r.exists]
        if not missing:
            continue
        if len(missing) == len(items):
            parts.append(f"缺失家族 {family}（两路皆无）")
        else:
            parts.append(f"家族 {family} 缺路次 {', '.join(missing)}")
    return "；".join(parts) if parts else ""


def build_asr_missing_error(batch: str, routes: list[RouteCheck]) -> PreflightError:
    inventory = format_route_inventory(routes)
    summary = summarize_missing(routes)
    message = (
        f"batch={batch} 分类前置检查失败，缺少以下 ASR 结果：\n"
        + "\n".join(inventory)
        + f"\n摘要: {summary}。\n"
        "请确认推理时 --source-name/--batch 与 --asr-run 是否与上表别名逐字一致；"
        "历史别名请用 --family 显式传入，勿依赖静默猜测。"
    )
    return PreflightError(batch, message, routes)


def check_asr_routes(batch: str, families: dict[str, list[str]], asr_dir: Path) -> list[RouteCheck]:
    routes: list[RouteCheck] = []
    for family, aliases in families.items():
        for alias in aliases:
            path = asr_dir / f"{alias}_asr_{batch}.parquet"
            routes.append(RouteCheck(family=family, alias=alias, path=path, exists=path.is_file()))
    return routes


def preflight(
    *,
    batch: str,
    cleaned: Path,
    families: dict[str, list[str]],
    asr_dir: Path,
    source_dir: Path | None,
) -> tuple[list, list[RouteCheck]]:
    """Read-only inventory before classification. Raises PreflightError / ValueError."""
    if source_dir is not None and not source_dir.is_dir():
        raise ValueError(
            f"原始音频目录不存在: {source_dir}（仅作存在性提示，不会重扫/重洗音频）"
        )
    if not cleaned.is_file():
        raise ValueError(f"cleaned 底表不存在: {cleaned}")

    base = list(Manifest.load(cleaned))
    validate_base_snapshot(base)
    if not base:
        raise ValueError(f"cleaned 底表为空: {cleaned}")
    missing_hash = [s.id for s in base if not original_audio_sha256(s)]
    if missing_hash:
        preview = ", ".join(missing_hash[:5])
        more = f" 等 {len(missing_hash)} 条" if len(missing_hash) > 5 else ""
        raise ValueError(f"cleaned 底表缺少 original_audio_sha256: {preview}{more}")

    unreadable = []
    for sample in base:
        path = sample.audio.get("resampled_16k", "__missing__")
        if not Path(path).is_file():
            unreadable.append(f"{sample.id}:{path}")
    if unreadable:
        preview = "; ".join(unreadable[:3])
        more = f" 等 {len(unreadable)} 条" if len(unreadable) > 3 else ""
        raise ValueError(
            f"resampled_16k 不可读: {preview}{more}；请恢复底表中的音频路径"
        )

    routes = check_asr_routes(batch, families, asr_dir)
    if any(not r.exists for r in routes):
        raise build_asr_missing_error(batch, routes)
    return base, routes


def validate_delivery(base, classified):
    expected = {(s.id, original_audio_sha256(s)) for s in base}
    actual = [(s.id, original_audio_sha256(s)) for s in classified]
    if len(actual) != len(expected) or set(actual) != expected:
        raise ValueError("分类前后样本 ID/原音频哈希不守恒")
    for sample in classified:
        labels = sample.labels
        if not labels.get("classification_bucket"):
            raise ValueError(f"缺少分类桶: {sample.id}")
        if labels.get("outcome") == "excluded":
            if not (labels.get("classification_reason_codes") or labels.get("exclude_reason")):
                raise ValueError(f"排除样本缺少原因: {sample.id}")
        elif labels.get("category") not in CATEGORIES:
            raise ValueError(f"未完成五分类: {sample.id}: {labels.get('category')}")


def _append_summary_stats(summary, samples, *, engine: str) -> None:
    summary.append(["category / outcome", "数量"])
    counts = Counter(s.labels.get("category") or s.labels.get("outcome") for s in samples)
    for name, count in sorted(counts.items()):
        summary.append([name, count])
    summary.append(["总计", len(samples)])
    if engine == "v2_3":
        vad_failed = sum(
            1
            for s in samples
            if str((s.quality or {}).get("vad_speech_state") or "").lower() == "failed"
            or str((s.quality or {}).get("vad_status") or "").lower() == "failed"
        )
        eat_required = 0
        for s in samples:
            cand = (s.labels or {}).get("efficientat_v2_3_candidate")
            if isinstance(cand, dict) and cand.get("required"):
                eat_required += 1
            elif cand is True:
                eat_required += 1
        acoustic_counts = Counter(
            str((s.quality or {}).get("acoustic_state") or (s.labels or {}).get("acoustic_state") or "")
            for s in samples
        )
        summary.append(["VAD failed", vad_failed])
        summary.append(["EfficientAT 调用/有状态条数", eat_required])
        summary.append(["acoustic_state 分布", str(dict(sorted(acoustic_counts.items())))])
        summary.append(
            [
                "说明",
                "v2.3：主声学证据为 Silero+EfficientAT；DNSMOS 不参与类别决策；"
                "gold_candidate 为候选金标；excluded 保留排除原因",
            ]
        )
    else:
        summary.append(
            ["DNSMOS 回退 v2", sum(bool(s.labels.get("v2_fallback")) for s in samples)]
        )
        summary.append(
            [
                "说明",
                "legacy v2.2：gold_candidate 为候选金标；excluded 保留排除原因；不代表人工金标发布",
            ]
        )


def export_workbook(samples, work_dir, output, *, engine: str = "v2_3"):
    # Reuse the authoritative export column contract, then add audit counts.
    from openpyxl import load_workbook

    from audio_engine.cli.main import review_export_summary

    if len(samples) > 1_048_575:
        raise ValueError("超过单表 Excel 行数上限 1,048,575，请拆分 batch")
    classified_path = work_dir / "classified.parquet"
    Manifest(samples).save(classified_path)
    temporary = work_dir / "summary.xlsx"
    review_export_summary(
        dataset=str(classified_path),
        output=temporary,
        output_manifest=None,
        max_rows=1_048_575,
        catalog_dir=ROOT / "data/catalog",
    )
    book = load_workbook(temporary)
    sheet = book.active
    sheet.title = "分类明细"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    # ASR text is data, including text beginning with '='; never Excel formulas.
    for row in sheet:
        for cell in row:
            if cell.data_type == "f":
                cell.data_type = "s"
    headers = {cell.value: cell.column for cell in sheet[1]}
    exported_ids = [sheet.cell(i, headers["sample_id"]).value for i in range(2, sheet.max_row + 1)]
    if exported_ids != [s.id for s in samples]:
        raise ValueError("XLSX 导出样本对账失败")
    summary = book.create_sheet("分类统计", 0)
    _append_summary_stats(summary, samples, engine=engine)
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 90
    book.save(temporary)
    book.close()
    os.replace(temporary, output)


def default_xlsx_path(batch: str, profile: dict) -> Path:
    return ROOT / "data" / "exports" / f"{profile['xlsx_stem']}_{batch}.xlsx"


def default_classified_path(batch: str, profile: dict) -> Path:
    return (
        ROOT
        / "datasets"
        / "stage1"
        / "derived"
        / f"{profile['classified_stem']}_{batch}.parquet"
    )


def atomic_save_manifest(samples, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="classified_", dir=output.parent) as folder:
        temporary = Path(folder) / "classified.parquet"
        Manifest(samples).save(temporary)
        os.replace(temporary, output)


def attach_audio_energy(samples, energy_cfg: OperatorConfig, workers: int):
    """CPU/IO-bound energy step; threads help when many cores and local disk."""
    workers = max(1, int(workers))
    energy = AudioEnergyOperator()
    total = len(samples)
    if workers == 1 or total <= 1:
        for i, sample in enumerate(samples):
            samples[i] = energy.process(sample, energy_cfg).sample
            if (i + 1) % 1000 == 0 or (i + 1) == total:
                print(f"  energy {i + 1}/{total} (workers=1)", flush=True)
        return samples

    print(f"  energy workers={workers}", flush=True)
    out = [None] * total
    done = 0
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {
            pool.submit(energy.process, sample, energy_cfg): idx
            for idx, sample in enumerate(samples)
        }
        for fut in as_completed(futures):
            idx = futures[fut]
            out[idx] = fut.result().sample
            done += 1
            if done % 1000 == 0 or done == total:
                print(f"  energy {done}/{total} (workers={workers})", flush=True)
    return out


def attach_silero_vad(samples, vad_cfg: OperatorConfig):
    """Full-pass Silero; failures stay auditable (never forge present/absent)."""
    vad = SileroVadOperator()
    total = len(samples)
    for i, sample in enumerate(samples):
        samples[i] = vad.process(sample, vad_cfg).sample
        if (i + 1) % 1000 == 0 or (i + 1) == total:
            print(f"  silero_vad {i + 1}/{total}", flush=True)
    return samples


def run_v23_engine(samples, *, cfg, cfg_path: Path, batch: str):
    """Match pipelines/classify_dataset_five_class_v2_3_min_acoustic.yaml step order."""
    print("[3/5] Silero VAD（全量）", flush=True)
    samples = attach_silero_vad(
        samples, cfg(config_path="configs/quality/silero_vad_v1.yaml")
    )
    print("[4/5] EfficientAT 候选 + 评分 → v2.3 五分类", flush=True)
    samples = EfficientAtCandidatesOperator().run(
        samples,
        cfg(
            config_path=str(cfg_path),
            selection_config=str(cfg_path),
            efficientat_config="configs/quality/efficientat_v1.yaml",
        ),
    )
    samples = EfficientAtOperator().run(
        samples,
        cfg(
            config_path=str(cfg_path),
            selection_config=str(cfg_path),
            efficientat_config="configs/quality/efficientat_v1.yaml",
            acoustic_evidence_config="configs/quality/acoustic_evidence_v1.yaml",
        ),
    )
    samples = ClassifyOperator().run(
        samples,
        cfg(
            config_path=str(cfg_path),
            rule_version=RULE_VERSION_FIVE_CLASS_V2_3,
            noise_policy="asr_anomaly_noise_v1",
            classify_text_policy="five_class_v1",
            quality={"calibrated": False, "noise_policy": "asr_anomaly_noise_v1"},
        ),
    )
    return samples


def run_v22_engine(samples, *, cfg, cfg_path: Path, batch: str, dnsmos_config: Path):
    print("[3/4] DNSMOS 候选评分和 v2.2 五分类", flush=True)
    sidecar = ROOT / f"datasets/stage1/derived/quality_sidecar_{batch}.parquet"
    samples = DnsmosV22CandidatesOperator().run(
        samples,
        cfg(
            config_path=str(cfg_path),
            dnsmos_config=str(dnsmos_config),
            quality_sidecar_manifest=str(sidecar) if sidecar.is_file() else None,
        ),
    )
    samples = ClassifyOperator().run(samples, cfg(config_path=str(cfg_path)))
    return samples


def run(args):
    validate_source_name(args.batch)
    profile = resolve_rule_profile(args.rule_version)
    if profile["engine"] == "v2_3":
        assert_v23_dependencies_ready()

    families, family_source = resolve_families(args)
    cleaned = args.cleaned or ROOT / f"datasets/stage1/cleaned/cleaned_{args.batch}.parquet"
    base, routes = preflight(
        batch=args.batch,
        cleaned=cleaned,
        families=families,
        asr_dir=args.asr_dir,
        source_dir=args.source_dir,
    )
    for sample in base:
        # source_path/raw preserve historical provenance, not the current audio
        # location. Post-ASR quality operators consume resampled_16k only.
        # The audio base must not contribute stale ASR or certification labels.
        sample.transcripts = {}
        sample.labels.pop("run_identities_verified", None)

    output = args.output or default_xlsx_path(args.batch, profile)
    if output.suffix.lower() != ".xlsx":
        raise ValueError("--output 必须使用 .xlsx 后缀")
    classified_out = None
    if not args.no_classified_parquet:
        classified_out = args.classified_output or default_classified_path(args.batch, profile)
        if classified_out.suffix.lower() != ".parquet":
            raise ValueError("--classified-output 必须使用 .parquet 后缀")

    existing = []
    if output.exists() and not args.overwrite:
        existing.append(str(output))
    if classified_out is not None and classified_out.exists() and not args.overwrite:
        existing.append(str(classified_out))
    if existing:
        raise ValueError(
            "输出已存在: " + "; ".join(existing) + "；如需替换请指定 --overwrite"
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    route_lines = "\n".join(format_route_inventory(routes))
    steps = 5 if profile["engine"] == "v2_3" else 4
    print(
        f"[0/{steps}] 前置检查通过 batch={args.batch} "
        f"rule_version={args.rule_version} engine={profile['engine']} "
        f"别名来源={family_source}\n{route_lines}",
        flush=True,
    )

    # Keep temporary files beside the final workbook so publication is atomic.
    with tempfile.TemporaryDirectory(prefix=f"classify_{args.batch}_", dir=output.parent) as folder:
        work = Path(folder)
        selection = load_yaml(ROOT / profile["selection_config"])
        selection.update(
            model_families=families,
            teacher_families=[f for f in families if f != "qwen"],
            run_aliases={a: f for f, aliases in families.items() for a in aliases},
        )
        # Keep rule_version explicit so classify cannot drift.
        selection["rule_version"] = profile["selection_rule"]
        selection["policy_version"] = profile["selection_rule"]
        SelectionV3Config.from_params(selection)
        cfg_path = work / "selection.yaml"
        cfg_path.write_text(yaml.safe_dump(selection, allow_unicode=True), encoding="utf-8")

        def cfg(**params):
            return OperatorConfig(
                params=params,
                run_dir=work,
                cache_dir=work / "cache",
                force=True,
                step_name="post_asr",
            )

        print(f"[1/{steps}] 对齐六路 ASR，共 {len(base)} 条音频", flush=True)
        samples = AggregateManifestsOperator().run(
            base,
            cfg(
                manifests=[{"model": r.alias, "path": str(r.path)} for r in routes],
                id_policy="exact",
                hash_policy="original_audio",
                require_hashes=True,
            ),
        )
        print(f"[2/{steps}] 计算音频能量", flush=True)
        energy_cfg = cfg(
            config_path="configs/quality/audio_energy_v1.yaml", selection_config=str(cfg_path)
        )
        samples = attach_audio_energy(samples, energy_cfg, workers=args.energy_workers)

        if profile["engine"] == "v2_3":
            samples = run_v23_engine(samples, cfg=cfg, cfg_path=cfg_path, batch=args.batch)
        else:
            samples = run_v22_engine(
                samples,
                cfg=cfg,
                cfg_path=cfg_path,
                batch=args.batch,
                dnsmos_config=args.dnsmos_config,
            )

        validate_delivery(base, samples)
        samples.sort(key=lambda s: (s.labels.get("category") or "excluded", s.id))
        print(f"[{steps}/{steps}] 导出并核对 XLSX", flush=True)
        export_workbook(samples, work, output, engine=profile["engine"])
        if classified_out is not None:
            atomic_save_manifest(samples, classified_out)
            print(f"分类中间表: {classified_out}", flush=True)
    print(f"完成: {output}", flush=True)
    return output


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=__doc__,
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--batch", required=True, help="与各家族推理时同一 batch / source-name")
    parser.add_argument(
        "--rule-version",
        default=DEFAULT_RULE_VERSION,
        choices=sorted(RULE_PROFILES.keys()),
        help=(
            f"分拣规则短名；默认 {RULE_CLI_V23}（034 最小声学）；"
            f"{RULE_CLI_V22} 为 legacy DNSMOS 口径"
        ),
    )
    parser.add_argument(
        "--source-dir",
        "--source_dir",
        type=Path,
        default=None,
        help="可选；仅做目录存在性提示，不会重扫/重洗音频。分类读底表 resampled_16k",
    )
    parser.add_argument(
        "--family",
        action="append",
        help="例如 qwen=qwen_1,qwen_2；可重复三次。优先于 dataset YAML 与 stage1 默认别名",
    )
    parser.add_argument(
        "--dataset-config",
        type=Path,
        help="显式批次 YAML（只读 model_families）；未指定时尝试 zh_asr_v3_<batch>.yaml",
    )
    parser.add_argument("--cleaned", type=Path, help="完整 cleaned 音频底表路径")
    parser.add_argument("--asr-dir", type=Path, default=ROOT / "datasets/stage1/asr")
    parser.add_argument(
        "--dnsmos-config",
        type=Path,
        default=ROOT / "configs/quality/dnsmos_p835.yaml",
        help=f"仅 --rule-version {RULE_CLI_V22} 使用；v2.3 默认不跑 DNSMOS",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="默认 data/exports/<stem>_<batch>.xlsx（stem 随 --rule-version）",
    )
    parser.add_argument(
        "--classified-output",
        type=Path,
        help="默认 datasets/stage1/derived/<stem>_<batch>.parquet（stem 随 --rule-version）",
    )
    parser.add_argument(
        "--no-classified-parquet",
        action="store_true",
        help="仅交付 XLSX，不写正式 classified parquet",
    )
    parser.add_argument(
        "--energy-workers",
        "--workers",
        type=int,
        default=1,
        dest="energy_workers",
        help="仅加速音频能量步线程并发（默认 1；服务器建议 16~32）。不控制 Silero/EfficientAT",
    )
    parser.add_argument("--overwrite", action="store_true")
    return parser


def main(argv=None):
    parser = build_parser()
    args = parser.parse_args(argv)
    # Resolve user paths before switching to the repository for config references.
    for key, value in vars(args).items():
        if isinstance(value, Path):
            setattr(args, key, value.resolve())
    previous = Path.cwd()
    try:
        os.chdir(ROOT)
        run(args)
    except PreflightError as exc:
        parser.exit(2, f"错误: {exc}\n")
    except (ValueError, TypeError, FileNotFoundError, KeyError) as exc:
        parser.exit(2, f"错误: {exc}\n")
    finally:
        os.chdir(previous)


if __name__ == "__main__":
    main()
