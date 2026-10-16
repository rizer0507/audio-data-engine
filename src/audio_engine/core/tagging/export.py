"""Multi-tag export: parquet labels, rule details, XLSX summary, filter export."""

from __future__ import annotations

import json
import math
from collections import Counter
from pathlib import Path

from audio_engine.core.sample import Sample
from audio_engine.core.tagging.engine import SampleTaggingResult
from audio_engine.core.tagging.types import POLICY_VERSION


def dnsmos_sig_for_reference(sample: Sample) -> float | None:
    """Return an available SIG score, never a stale score from an unsuccessful run."""
    quality = sample.quality if isinstance(sample.quality, dict) else {}
    if str(quality.get("dnsmos_status") or "").strip().lower() != "success":
        return None
    value = quality.get("dnsmos_sig")
    if value is None or isinstance(value, bool):
        return None
    try:
        score = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return score if math.isfinite(score) else None


def sort_samples_by_dnsmos_sig(samples: list[Sample]) -> list[Sample]:
    """Stable ascending SIG reference order; retain samples without available scores."""

    def sort_key(sample: Sample) -> tuple[bool, float]:
        score = dnsmos_sig_for_reference(sample)
        return score is None, score if score is not None else 0.0

    return sorted(samples, key=sort_key)


def apply_result_to_sample(
    sample: Sample,
    result: SampleTaggingResult,
    *,
    policy_version: str = POLICY_VERSION,
    snapshot_id: str = "",
) -> Sample:
    labels = dict(sample.labels or {})
    labels["tags"] = list(result.effective_tags)
    labels["tagging"] = {
        "policy_version": policy_version,
        "snapshot_id": snapshot_id,
        "eval_status": result.eval_status,
        "asr_fingerprint": result.asr_fingerprint,
        "untagged_reason": result.untagged_reason,
        "rule_coverage": {
            rr.rule_id: rr.ternary.value for rr in result.rule_results
        },
    }
    sample.labels = labels
    return sample


def write_rule_details(
    path: Path, results, *, batch_id: str, snapshot_id: str, policy_version: str,
) -> Path:
    """Stream audit rows; never accumulate all rules from a million-sample run."""
    from contextlib import ExitStack

    from audio_engine.core.artifacts import atomic_path
    path = Path(path)
    with atomic_path(path) as tmp, ExitStack() as stack:
        stream = stack.enter_context(tmp.open("w", encoding="utf-8")) if path.suffix.lower() != ".parquet" else None
        writer = None
        rows = []
        for result in results:
            for rr in result.rule_results:
                row = {"batch_id": batch_id, "snapshot_id": snapshot_id,
                       "policy_version": policy_version, "sample_id": result.sample_id,
                       "rule_id": rr.rule_id, "rule_version": rr.rule_version, "tag_id": rr.tag_id,
                       "ternary": rr.ternary.value,
                       "acoustic_plan": json.dumps(rr.acoustic_plan, ensure_ascii=False),
                       "condition": json.dumps(rr.condition.as_dict(), ensure_ascii=False)}
                if stream is not None:
                    stream.write(json.dumps(row, ensure_ascii=False) + "\n")
                else:
                    rows.append(row)
                    if len(rows) >= 1000:
                        import pyarrow as pa
                        import pyarrow.parquet as pq
                        table = pa.Table.from_pylist(rows)
                        if writer is None:
                            writer = stack.enter_context(pq.ParquetWriter(tmp, table.schema))
                        writer.write_table(table)
                        rows.clear()
        if stream is None:
            import pyarrow as pa
            import pyarrow.parquet as pq
            if writer is not None:
                if rows:
                    writer.write_table(pa.Table.from_pylist(rows, schema=writer.schema))
            else:
                pq.write_table(pa.Table.from_pylist(rows), tmp)
    return path


def write_summary_xlsx(
    path: Path, samples: list[Sample], results, *, batch_id: str,
    overlap_tags: list[str] | None = None, ordered: bool = False,
) -> Path:
    """Write-only XLSX preserves the existing sheets without retaining every row."""
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell

    from audio_engine.core.artifacts import atomic_path
    book = Workbook(write_only=True)
    sample_sheet = book.create_sheet("samples")
    sample_sheet.append(["sample_id", "sha256", "tags", "tag_count", "eval_status",
                         "untagged_reason", "dnsmos_status", "dnsmos_sig"])
    tag_counter, status_counter, failed_deps = Counter(), Counter(), Counter()
    unknown_rules = 0
    # Public callers may supply results out of order; the service explicitly guarantees order.
    by_id = None if ordered else {r.sample_id: r for r in results}
    iterator = iter(results) if ordered else None
    overlap = book.create_sheet("overlap") if overlap_tags else None
    if overlap:
        overlap.append(["sample_id", "tags"])
    wanted = set(overlap_tags or [])
    for sample in samples:
        result = next(iterator, None) if ordered else by_id.get(sample.id)
        if result is not None and result.sample_id != sample.id:
            raise ValueError("ordered tagging summary received mismatched sample IDs")
        tags = list(result.effective_tags) if result else list((sample.labels or {}).get("tags") or [])
        status = result.eval_status if result else str((sample.labels.get("tagging") or {}).get("eval_status") or "")
        status_counter[status or "unknown"] += 1
        tag_counter.update(tags)
        if result:
            for rr in result.rule_results:
                unknown_rules += rr.ternary.value == "unknown"
                failed_deps.update(rr.condition.missing_deps)
        sig_cell = WriteOnlyCell(sample_sheet, value=dnsmos_sig_for_reference(sample))
        sig_cell.number_format = "0.000"
        sample_sheet.append([sample.id, sample.sha256, ",".join(tags), len(tags), status,
                             result.untagged_reason if result else None,
                             (sample.quality or {}).get("dnsmos_status"), sig_cell])
        if overlap and wanted.issubset(tags):
            overlap.append([sample.id, ",".join(sorted(tags))])
    for name, headers, rows in (
        ("tag_counts", ["tag_id", "sample_count"], sorted(tag_counter.items())),
        ("eval_status", ["eval_status", "count"], sorted(status_counter.items())),
        ("overview", ["metric", "value"], [("sample_count", len(samples)),
         ("tag_assignment_count", sum(tag_counter.values())), ("unknown_rule_judgements", unknown_rules),
         ("status_sum", sum(status_counter.values()))]),
        ("failed_deps", ["dependency", "count"], sorted(failed_deps.items())),
    ):
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
    with atomic_path(Path(path)) as tmp:
        book.save(tmp)
    return Path(path)


def write_tagged_parquet(path: Path, samples: list[Sample], *, batch_size: int = 1000) -> Path:
    """Retain flattened columns while bounding export buffers and handling null-first blocks."""
    import numbers

    import pyarrow as pa
    import pyarrow.parquet as pq

    from audio_engine.core.artifacts import atomic_path
    kinds = {}
    for sample in samples:
        for key, value in sample.to_flat_dict().items():
            observed = kinds.setdefault(key, set())
            if value is not None:
                observed.add("bool" if isinstance(value, bool) else "int" if isinstance(value, numbers.Integral)
                             else "float" if isinstance(value, numbers.Real) else "str")
    types = {key: pa.bool_() if values == {"bool"} else pa.int64() if values == {"int"}
             else pa.float64() if values and values <= {"int", "float"} else pa.string()
             for key, values in kinds.items()}
    schema = pa.schema(list(types.items()))
    with atomic_path(Path(path)) as tmp, pq.ParquetWriter(tmp, schema) as writer:
        for start in range(0, len(samples), batch_size):
            rows = []
            for sample in samples[start:start + batch_size]:
                row = sample.to_flat_dict()
                for key, value in row.items():
                    if types[key] == pa.string() and value is not None and not isinstance(value, str):
                        row[key] = json.dumps(value, ensure_ascii=False, default=str)
                rows.append(row)
            writer.write_table(pa.Table.from_pylist(rows, schema=schema))
    return Path(path)


def filter_samples(
    samples: list[Sample],
    *,
    include_all: list[str] | None = None,
    include_any: list[str] | None = None,
    exclude_tags: list[str] | None = None,
    require_unknown: bool = False,
    require_manual_block: bool = False,
) -> list[Sample]:
    out: list[Sample] = []
    for sample in samples:
        labels = sample.labels if isinstance(sample.labels, dict) else {}
        tags = set(labels.get("tags") or [])
        tagging = labels.get("tagging") if isinstance(labels.get("tagging"), dict) else {}
        if include_all and not set(include_all).issubset(tags):
            continue
        if include_any and not (set(include_any) & tags):
            continue
        if exclude_tags and (set(exclude_tags) & tags):
            continue
        if require_unknown:
            coverage = tagging.get("rule_coverage") or {}
            if not any(v == "unknown" for v in coverage.values()):
                continue
        if require_manual_block:
            # Best-effort: untagged_reason or sources not always on sample; check tagging note.
            if tagging.get("untagged_reason") != "matched_but_manually_blocked":
                # Also allow if tags empty but we cannot see blocks — skip unless reason set.
                continue
        out.append(sample)
    return out
