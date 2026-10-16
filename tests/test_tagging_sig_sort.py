"""SIG is an optional export reference; it cannot remove review candidates."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from audio_engine.core.manifest import Manifest
from audio_engine.core.sample import Sample
from audio_engine.core.tagging.config import TaggingConfig
from audio_engine.core.tagging.evidence_config import TagEvidenceConfig
from audio_engine.core.tagging.export import sort_samples_by_dnsmos_sig, write_summary_xlsx
from audio_engine.core.tagging.service import TaggingService


def _sample(sample_id: str, quality: dict | None = None, *, tags: list[str] | None = None) -> Sample:
    sample = Sample(
        id=sample_id,
        source_path=f"/audio/{sample_id}.wav",
        sha256="a" * 64,
        sample_rate=16000,
        channels=1,
        duration=3.5,
        audio={"raw": f"/audio/{sample_id}.wav", "normalized": f"/norm/{sample_id}.wav"},
        quality=quality or {},
        transcripts={"qwen_1": {"text": "可能是这句话", "extra": {"run": 1}}},
        labels={
            "tags": ["review"] if tags is None else tags,
            "parent_sample_id": "parent",
            "segment_index": 2,
            "tagging": {"eval_status": "tagged", "rule_coverage": {"review_rule": "matched"}},
        },
        status={"qwen_1": "completed"},
        errors={"other_run": "test error"},
    )
    sample.add_lineage("segment", "1", {"offset": 1.0}, input_key="raw", output_key="raw")
    return sample


def _service() -> TaggingService:
    return TaggingService(
        TaggingConfig(
            policy_version="test",
            evidence=TagEvidenceConfig(),
            evidence_mode="existing",
            tags={},
            rules=[],
            fragments={},
            lexicons={},
        )
    )


def _payload(sample: Sample) -> str:
    # JSON's NaN/Infinity representation also lets us compare untouched non-finite metadata.
    return json.dumps(sample.model_dump(), sort_keys=True, ensure_ascii=False)


def test_sig_sort_stably_retains_unavailable_scores_and_ignores_failed_residue():
    samples = [
        _sample("missing"),
        _sample("high", {"dnsmos_status": "success", "dnsmos_sig": 3.8}),
        _sample("failed", {"dnsmos_status": "failed", "dnsmos_sig": 0.1}),
        _sample("low-first", {"dnsmos_status": "success", "dnsmos_sig": 1.7}),
        _sample("pending", {"dnsmos_status": "not_computed", "dnsmos_sig": 0.2}),
        _sample("low-tie", {"dnsmos_status": "success", "dnsmos_sig": 1.7}),
        _sample("invalid", {"dnsmos_status": "success", "dnsmos_sig": "invalid"}),
        _sample("nan", {"dnsmos_status": "success", "dnsmos_sig": float("nan")}),
        _sample("positive-inf", {"dnsmos_status": "success", "dnsmos_sig": float("inf")}),
        _sample("negative-inf", {"dnsmos_status": "success", "dnsmos_sig": -float("inf")}),
        _sample("null-score", {"dnsmos_status": "success", "dnsmos_sig": None}),
        _sample("no-status", {"dnsmos_sig": 0.3}),
        _sample("bool-score", {"dnsmos_status": "success", "dnsmos_sig": False}),
        _sample("numeric-string", {"dnsmos_status": " SUCCESS ", "dnsmos_sig": "2.4"}),
    ]
    before = [_payload(sample) for sample in samples]

    ordered = sort_samples_by_dnsmos_sig(samples)

    assert [sample.id for sample in ordered] == [
        "low-first", "low-tie", "numeric-string", "high", "missing", "failed", "pending",
        "invalid", "nan", "positive-inf", "negative-inf", "null-score", "no-status", "bool-score",
    ]
    assert len(ordered) == len(samples)
    assert all(any(result is original for original in samples) for result in ordered)
    assert [_payload(sample) for sample in samples] == before


def test_export_default_keeps_input_order_and_full_sample_payload(tmp_path: Path):
    samples = [
        _sample("missing"),
        _sample("high", {"dnsmos_status": "success", "dnsmos_sig": 4.0}),
        _sample("failed", {"dnsmos_status": "failed", "dnsmos_sig": 0.5}),
        _sample("low", {"dnsmos_status": "success", "dnsmos_sig": 1.8}),
    ]
    output = tmp_path / "default.parquet"

    result = _service().export_filtered(samples, output=output)
    restored = list(Manifest.load(output))

    assert result["count"] == len(samples)
    assert [_payload(sample) for sample in restored] == [_payload(sample) for sample in samples]


def test_export_sorts_filtered_manifest_and_retains_missing_nonfinite_metadata(tmp_path: Path):
    samples = [
        _sample("missing"),
        _sample("high", {"dnsmos_status": "success", "dnsmos_sig": 4.0}),
        _sample("failed", {"dnsmos_status": "failed", "dnsmos_sig": 0.1}),
        _sample("low-first", {"dnsmos_status": "success", "dnsmos_sig": 1.8}),
        _sample("excluded", {"dnsmos_status": "success", "dnsmos_sig": 0.5}, tags=["other"]),
        _sample("nan", {"dnsmos_status": "success", "dnsmos_sig": float("nan")}),
        _sample("low-tie", {"dnsmos_status": "success", "dnsmos_sig": 1.8}),
        _sample("positive-inf", {"dnsmos_status": "success", "dnsmos_sig": float("inf")}),
        _sample("negative-inf", {"dnsmos_status": "success", "dnsmos_sig": -float("inf")}),
        _sample("blocked", {"dnsmos_status": "success", "dnsmos_sig": 0.6}, tags=["review", "block"]),
    ]
    output = tmp_path / "sorted.parquet"
    before = [_payload(sample) for sample in samples]

    result = _service().export_filtered(
        samples,
        output=output,
        include_all=["review"],
        exclude_tags=["block"],
        sort_by_dnsmos_sig=True,
    )
    restored = list(Manifest.load(output))

    assert [sample.id for sample in restored] == [
        "low-first", "low-tie", "high", "missing", "failed", "nan", "positive-inf", "negative-inf",
    ]
    assert result["count"] == len(samples) - 2
    original_by_id = {sample.id: sample for sample in samples}
    assert [_payload(sample) for sample in restored] == [
        _payload(original_by_id[sample.id]) for sample in restored
    ]
    assert [_payload(sample) for sample in samples] == before


def test_summary_xlsx_shows_only_available_sig_scores_and_preserves_all_samples(tmp_path: Path):
    samples = [
        _sample("valid", {"dnsmos_status": "success", "dnsmos_sig": 2.1}),
        _sample("failed", {"dnsmos_status": "failed", "dnsmos_sig": 0.1}),
        _sample("missing"),
        _sample("invalid", {"dnsmos_status": "success", "dnsmos_sig": "bad"}),
        _sample("nan", {"dnsmos_status": "success", "dnsmos_sig": float("nan")}),
        _sample("inf", {"dnsmos_status": "success", "dnsmos_sig": float("inf")}),
        _sample("skipped", {"dnsmos_status": "skipped_not_required", "dnsmos_sig": 0.2}),
    ]
    path = tmp_path / "summary.xlsx"

    write_summary_xlsx(path, samples, [], batch_id="test")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook["samples"].values)
        records = [dict(zip(rows[0], row)) for row in rows[1:]]
        assert [row["sample_id"] for row in records] == [sample.id for sample in samples]
        assert [row["dnsmos_sig"] for row in records] == [2.1, None, None, None, None, None, None]
        assert [row["dnsmos_status"] for row in records] == [
            "success", "failed", None, "success", "success", "success", "skipped_not_required",
        ]
        assert set(workbook.sheetnames) == {
            "samples", "tag_counts", "eval_status", "overview", "failed_deps",
        }
    finally:
        workbook.close()


@pytest.mark.parametrize("sort_flag", [False, True])
def test_tags_export_cli_passes_optional_sig_sort(tmp_path: Path, monkeypatch, sort_flag: bool):
    from typer.testing import CliRunner

    from audio_engine.cli import main as cli

    service = _service()
    monkeypatch.setattr(cli, "_tagging_service", lambda _config: service)
    samples = [
        _sample("missing"),
        _sample("high", {"dnsmos_status": "success", "dnsmos_sig": 4.0}),
        _sample("low", {"dnsmos_status": "success", "dnsmos_sig": 1.8}),
    ]
    input_path = tmp_path / "input.parquet"
    output_path = tmp_path / "output.parquet"
    Manifest(samples).save(input_path)
    args = ["tags", "export", "--input", str(input_path), "--output", str(output_path)]
    if sort_flag:
        args.append("--sort-by-dnsmos-sig")

    result = CliRunner().invoke(cli.app, args)

    assert result.exit_code == 0, result.output
    assert [sample.id for sample in Manifest.load(output_path)] == (
        ["low", "high", "missing"] if sort_flag else ["missing", "high", "low"]
    )
