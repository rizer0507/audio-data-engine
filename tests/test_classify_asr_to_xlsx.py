"""Real post-ASR integration without inference services or DNSMOS downloads."""

import argparse
import importlib.util
import sys
from pathlib import Path

import numpy as np
import pytest
import soundfile as sf
from openpyxl import load_workbook

from audio_engine.core.manifest import Manifest, file_sha256
from audio_engine.core.sample import Sample

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("post_asr", ROOT / "scripts/classify_asr_to_xlsx.py")
script = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = script
spec.loader.exec_module(script)


def _ns(**kwargs):
    defaults = dict(
        dataset_config=None,
        dnsmos_config=ROOT / "configs/quality/dnsmos_p835.yaml",
        overwrite=False,
        no_classified_parquet=True,
        classified_output=None,
        source_dir=None,
        family=None,
        cleaned=None,
        asr_dir=ROOT / "datasets/stage1/asr",
        output=None,
        energy_workers=1,
    )
    defaults.update(kwargs)
    return argparse.Namespace(**defaults)


@pytest.fixture
def batch(tmp_path, monkeypatch):
    monkeypatch.chdir(ROOT)
    source = tmp_path / "audio"
    source.mkdir()
    asr = tmp_path / "asr"
    asr.mkdir()
    samples = []
    for sid in ["0001", "0002"]:
        path = source / f"{sid}.wav"
        sf.write(path, np.zeros(16000), 16000)
        samples.append(
            Sample(
                id=sid,
                source_path=str(path),
                sha256=file_sha256(path),
                audio={"resampled_16k": str(path)},
                duration=1.0,
            )
        )
    cleaned = tmp_path / "cleaned.parquet"
    Manifest(samples).save(cleaned)
    families = {f: [f"{f}-a", f"{f}-b"] for f in ["qwen", "glm", "sensevoice"]}
    for aliases in families.values():
        for alias in aliases:
            incoming = [s.model_copy(deep=True) for s in reversed(samples)]
            for sample in incoming:
                sample.transcripts[alias] = {"text": "", "status": "success"}
            Manifest(incoming).save(asr / f"{alias}_asr_demo.parquet")
    return _ns(
        batch="demo",
        source_dir=source,
        asr_dir=asr,
        cleaned=cleaned,
        family=[f"{f}={','.join(a)}" for f, a in families.items()],
        output=tmp_path / "out.xlsx",
        no_classified_parquet=True,
    )


def test_end_to_end_single_workbook(batch):
    script.run(batch)
    book = load_workbook(batch.output)
    assert book.sheetnames == ["分类统计", "分类明细"]
    sheet = book["分类明细"]
    records = list(sheet.values)
    rows = [dict(zip(records[0], row)) for row in records[1:]]
    assert [r["sample_id"] for r in rows] == ["0001", "0002"]
    assert {r["category"] for r in rows} == {"environment_noise"}
    assert {r["noise_kind"] for r in rows} == {"silence"}
    assert len([h for h in records[0] if h.endswith(("-a_text", "-b_text"))]) == 6
    book.close()
    assert list(batch.output.parent.glob("classify_demo_*")) == []
    with pytest.raises(ValueError, match="输出已存在"):
        script.run(batch)


def test_energy_workers_parallel_happy_path(batch):
    batch.energy_workers = 4
    script.run(batch)
    book = load_workbook(batch.output)
    assert book.sheetnames == ["分类统计", "分类明细"]
    assert book["分类明细"].max_row == 3
    book.close()


def test_end_to_end_writes_classified_parquet(batch, tmp_path):
    batch.no_classified_parquet = False
    batch.classified_output = tmp_path / "classified.parquet"
    script.run(batch)
    assert batch.classified_output.is_file()
    classified = list(Manifest.load(batch.classified_output))
    assert [s.id for s in classified] == ["0001", "0002"]
    assert all(s.labels.get("category") == "environment_noise" for s in classified)


def test_hash_mismatch_stops_before_export(batch):
    path = batch.asr_dir / "glm-a_asr_demo.parquet"
    samples = list(Manifest.load(path))
    samples[0].sha256 = "f" * 64
    Manifest(samples).save(path)
    with pytest.raises(ValueError, match="mismatch"):
        script.run(batch)
    assert not batch.output.exists()


def test_missing_route_stops_with_inventory(batch):
    (batch.asr_dir / "glm-a_asr_demo.parquet").unlink()
    with pytest.raises(script.PreflightError, match="分类前置检查失败") as caught:
        script.run(batch)
    err = caught.value
    assert err.batch == "demo"
    by_alias = {r.alias: r for r in err.routes}
    assert by_alias["glm-a"].exists is False
    assert by_alias["glm-b"].exists is True
    assert by_alias["qwen-a"].exists is True
    message = str(err)
    assert "[glm] glm-a" in message and "缺失" in message
    assert "摘要:" in message
    assert "家族 glm 缺路次 glm-a" in message
    assert "请确认推理时" in message
    assert not batch.output.exists()


def test_missing_whole_family_summary(batch):
    (batch.asr_dir / "sensevoice-a_asr_demo.parquet").unlink()
    (batch.asr_dir / "sensevoice-b_asr_demo.parquet").unlink()
    with pytest.raises(script.PreflightError, match="缺失家族 sensevoice") as caught:
        script.run(batch)
    message = str(caught.value)
    assert "两路皆无" in message
    assert "[sensevoice] sensevoice-a" in message
    assert "[sensevoice] sensevoice-b" in message


def test_wrong_alias_does_not_silently_use_other_files(batch):
    batch.family = [
        "qwen=qwen-wrong-1,qwen-wrong-2",
        "glm=glm-a,glm-b",
        "sensevoice=sensevoice-a,sensevoice-b",
    ]
    with pytest.raises(script.PreflightError, match="qwen-wrong-1") as caught:
        script.run(batch)
    message = str(caught.value)
    assert "缺失家族 qwen" in message
    assert "qwen-a" not in message or "qwen-wrong" in message
    assert not batch.output.exists()


def test_duplicate_alias_rejected(batch):
    batch.family = ["qwen=a,b", "glm=c,d", "sensevoice=a,f"]
    with pytest.raises(ValueError, match="互异"):
        script.resolve_families(batch)


def test_stage1_default_aliases_when_no_family_or_dataset(tmp_path, monkeypatch):
    monkeypatch.chdir(ROOT)
    args = _ns(batch="demo-batch", family=None, dataset_config=None)
    # Avoid picking up a real workspace dataset YAML for this batch name.
    monkeypatch.setattr(
        script,
        "default_dataset_config_path",
        lambda batch: tmp_path / "missing.yaml",
    )
    families, source = script.resolve_families(args)
    assert source == "stage1_default"
    assert families == {
        "qwen": ["qwen_1", "qwen_2"],
        "glm": ["glm_1", "glm_2"],
        "sensevoice": ["sensevoice_1", "sensevoice_2"],
    }


def test_cli_family_beats_dataset_yaml(tmp_path, monkeypatch):
    monkeypatch.chdir(ROOT)
    cfg = tmp_path / "dataset.yaml"
    cfg.write_text(
        "model_families:\n"
        "  qwen: [qwen_1, qwen_2]\n"
        "  glm: [glm_1, glm_2]\n"
        "  sensevoice: [sensevoice_1, sensevoice_2]\n",
        encoding="utf-8",
    )
    args = _ns(
        batch="demo",
        dataset_config=cfg,
        family=["qwen=a,b", "glm=c,d", "sensevoice=e,f"],
    )
    families, source = script.resolve_families(args)
    assert source == "cli"
    assert families["qwen"] == ["a", "b"]


def test_unreadable_audio_stops(batch):
    (batch.source_dir / "0001.wav").unlink()
    with pytest.raises(ValueError, match="不可读"):
        script.run(batch)
    assert not batch.output.exists()


def test_historical_source_outside_current_directory_is_allowed(batch):
    samples = list(Manifest.load(batch.cleaned))
    for sample in samples:
        sample.source_path = f"/historical/unavailable/pcm/{sample.id}.pcm"
        sample.audio["raw"] = sample.source_path
    Manifest(samples).save(batch.cleaned)
    script.run(batch)
    book = load_workbook(batch.output)
    values = list(book["分类明细"].values)
    source_column = values[0].index("source_path")
    category_column = values[0].index("category")
    assert all(row[source_column].startswith("/historical/") for row in values[1:])
    assert {row[category_column] for row in values[1:]} == {"environment_noise"}
    book.close()


def test_failed_route_is_not_empty_vote(batch):
    path = batch.asr_dir / "glm-a_asr_demo.parquet"
    samples = list(Manifest.load(path))
    for sample in samples:
        sample.transcripts["glm-a"] = {"text": "", "status": "failed"}
    Manifest(samples).save(path)
    script.run(batch)
    book = load_workbook(batch.output)
    values = list(book["分类明细"].values)
    category = values[0].index("category")
    assert {row[category] for row in values[1:]} == {"hardcase"}
    book.close()


def test_export_formula_text_is_literal(tmp_path):
    sample = Sample(
        id="001",
        source_path="audio.wav",
        sha256="a" * 64,
        transcripts={"qwen-a": {"text": "=1+1"}},
        labels={"classification_bucket": "hardcase", "category": "hardcase"},
    )
    output = tmp_path / "result.xlsx"
    script.export_workbook([sample], tmp_path, output)
    book = load_workbook(output)
    sheet = book["分类明细"]
    column = next(c.column for c in sheet[1] if c.value == "qwen-a_text")
    assert sheet.cell(2, column).value == "=1+1"
    assert sheet.cell(2, column).data_type == "s"
    book.close()


def test_source_dir_optional_when_cleaned_self_contained(batch):
    batch.source_dir = None
    script.run(batch)
    assert batch.output.is_file()
